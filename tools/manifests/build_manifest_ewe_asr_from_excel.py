"""Build an ASR manifest TSV from the provided Excel metadata.

Input (present in this repo):
  data/Ewe/selected transcribed audios.xlsx

Audio files (present in this repo):
  data/Ewe/audios/*.mp3

This script rebuilds a clean manifest with reproducible splits.

Output TSV columns:
  split, audio, text, speaker_id, locale

Split strategy (default): deterministic speaker-hash split to avoid speaker leakage.

Usage:
  .\.venv\Scripts\python.exe tools/manifests/build_manifest_ewe_asr_from_excel.py --out data/manifests/asr/manifest_ewe_asr.tsv

Dependency:
- openpyxl
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import os
from dataclasses import dataclass
from typing import Any

import openpyxl


def _repo_root() -> str:
    return os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))


@dataclass(frozen=True)
class Row:
    split: str
    audio: str
    text: str
    speaker_id: str
    locale: str


def _stable_bucket(value: str) -> float:
    h = hashlib.md5(value.encode("utf-8")).digest()
    n = int.from_bytes(h[:8], "big")
    return n / float(2**64)


def _assign_split_speaker(speaker_id: str, dev_frac: float, test_frac: float) -> str:
    x = _stable_bucket(speaker_id)
    if x < test_frac:
        return "test"
    if x < test_frac + dev_frac:
        return "dev"
    return "train"


def _clean_text(text: Any) -> str:
    if text is None:
        return ""
    s = str(text)
    s = s.replace("\u00A0", " ").replace("\r", " ").replace("\n", " ")
    return " ".join(s.split()).strip()


def build_manifest(
    in_xlsx: str,
    audio_dir: str,
    out_tsv: str,
    dev_frac: float = 0.10,
    test_frac: float = 0.10,
) -> int:
    repo_root = _repo_root()

    wb = openpyxl.load_workbook(in_xlsx, read_only=True)
    ws = wb.active

    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    headers = [str(h).strip() if h is not None else "" for h in header_row]
    idx = {name: i for i, name in enumerate(headers)}

    required = ["AUDIO_PATH", "Transcription", "SPEAKER_ID", "LOCALE", "Full Filename"]
    missing = [c for c in required if c not in idx]
    if missing:
        raise SystemExit(f"Missing required columns in xlsx: {missing} (found: {headers})")

    rows: list[Row] = []
    missing_audio = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        full_filename = row[idx["Full Filename"]]
        speaker_id = row[idx["SPEAKER_ID"]]
        locale = row[idx["LOCALE"]]
        text = row[idx["Transcription"]]

        if full_filename is None:
            continue

        audio_path = os.path.join(audio_dir, str(full_filename))
        audio_rel = os.path.relpath(os.path.abspath(audio_path), repo_root).replace("\\", "/")

        if not os.path.exists(audio_path):
            missing_audio += 1

        speaker_str = str(speaker_id).strip() if speaker_id is not None else ""
        split = _assign_split_speaker(speaker_str or "<empty>", dev_frac=dev_frac, test_frac=test_frac)

        rows.append(
            Row(
                split=split,
                audio=audio_rel,
                text=_clean_text(text),
                speaker_id=speaker_str,
                locale=str(locale).strip() if locale is not None else "",
            )
        )

    rows.sort(key=lambda r: (r.split, r.speaker_id, r.audio))

    os.makedirs(os.path.dirname(out_tsv) or ".", exist_ok=True)
    with open(out_tsv, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, delimiter="\t")
        w.writerow(["split", "audio", "text", "speaker_id", "locale"])
        for r in rows:
            w.writerow([r.split, r.audio, r.text, r.speaker_id, r.locale])

    print("xlsx:", in_xlsx)
    print("audio_dir:", audio_dir)
    print("wrote:", out_tsv)
    print("rows:", len(rows))
    print("missing_audio_files:", missing_audio)

    return 0


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument(
        "--in-xlsx",
        default=os.path.join("data", "Ewe", "selected transcribed audios.xlsx"),
        help="Input XLSX (default: data/Ewe/selected transcribed audios.xlsx)",
    )
    ap.add_argument(
        "--audio-dir",
        default=os.path.join("data", "Ewe", "audios"),
        help="Audio directory containing mp3 files (default: data/Ewe/audios)",
    )
    ap.add_argument(
        "--out",
        default=os.path.join("data", "manifests", "asr", "manifest_ewe_asr.tsv"),
        help="Output TSV (default: data/manifests/asr/manifest_ewe_asr.tsv)",
    )
    ap.add_argument("--dev-frac", type=float, default=0.10)
    ap.add_argument("--test-frac", type=float, default=0.10)
    args = ap.parse_args()

    return build_manifest(args.in_xlsx, args.audio_dir, args.out, dev_frac=args.dev_frac, test_frac=args.test_frac)


if __name__ == "__main__":
    raise SystemExit(main())
